from __future__ import annotations

import json
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import List

from openpyxl import Workbook, load_workbook


@dataclass
class SubmissionRecord:
    employee_id: str
    employee_name: str
    year: int
    quarter: str
    form_type: str
    answers: List[float]
    comments: List[str]
    submitted_at: str


class ExcelFeedbackStore:
    FEEDBACK_SHEET = "feedback"
    SUMMARY_VIEW_SHEET = "Sheet1"
    QUEUE_SHEET = "ingestion_queue"

    HEADER = [
        "record_type",  # submission | summary
        "employee_id",
        "employee_name",
        "year",
        "quarter",
        "form_type",
        "answers_json",
        "comments_json",
        "submitted_at",
        "manager_email",
        "client_email",
        "peer_1",
        "peer_2",
        "self_avg",
        "client_avg",
        "peer_avg",
        "manager_avg",
        "total_average",
        "client_diff_gt_1",
        "manager_diff_gt_1",
        "client_diff_gt_05",
        "manager_diff_gt_05",
    ]
    SUMMARY_VIEW_HEADER = [
        "Emp ID",
        "Name",
        "Manager Email",
        "Client Email",
        "Peer 1",
        "Peer 2",
        "Self",
        "Client",
        "Peer",
        "Manager",
        "Total Average",
        "Client Difference >1",
        "Manager Difference >1",
        "Client Difference >0.5",
        "Manager Difference >0.5",
    ]
    QUEUE_HEADER = [
        "event_id",
        "status",
        "attempt_count",
        "next_retry_at",
        "created_at",
        "updated_at",
        "last_error",
        "payload_json",
    ]

    def __init__(self, workbook_path: str, summary_view_sheet: str = SUMMARY_VIEW_SHEET) -> None:
        self.path = Path(workbook_path)
        self.summary_view_sheet = summary_view_sheet
        self._lock = Lock()
        self._ensure_workbook()

    def _ensure_workbook(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            wb = load_workbook(self.path)
            if self.FEEDBACK_SHEET not in wb.sheetnames:
                ws = wb.create_sheet(self.FEEDBACK_SHEET)
                ws.append(self.HEADER)
                wb.save(self.path)
                return
            ws = wb[self.FEEDBACK_SHEET]
            headers = [ws.cell(row=1, column=c).value for c in range(1, len(self.HEADER) + 1)]
            if headers != self.HEADER:
                raise ValueError(
                    "Workbook header mismatch for single-sheet mode. "
                    "Use a fresh workbook path or migrate existing sheet."
                )
            self._ensure_summary_view_sheet(wb)
            self._ensure_queue_sheet(wb)
            wb.save(self.path)
            return

        wb = Workbook()
        ws = wb.active
        ws.title = self.FEEDBACK_SHEET
        ws.append(self.HEADER)
        wb.create_sheet(self.summary_view_sheet).append(self.SUMMARY_VIEW_HEADER)
        wb.create_sheet(self.QUEUE_SHEET).append(self.QUEUE_HEADER)
        wb.save(self.path)

    def _normalize_headers(self, headers: List[str | None]) -> List[str]:
        return [str(h or "").strip() for h in headers]

    def _ensure_summary_view_sheet(self, wb) -> None:
        if self.summary_view_sheet not in wb.sheetnames:
            ws = wb.create_sheet(self.summary_view_sheet)
            ws.append(self.SUMMARY_VIEW_HEADER)
            return
        ws = wb[self.summary_view_sheet]
        existing = [ws.cell(row=1, column=c).value for c in range(1, len(self.SUMMARY_VIEW_HEADER) + 1)]
        if self._normalize_headers(existing) != self._normalize_headers(self.SUMMARY_VIEW_HEADER):
            raise ValueError(
                f"Summary sheet '{self.summary_view_sheet}' header mismatch. "
                "Please align the first 15 columns with the expected HR table."
            )

    def _ensure_queue_sheet(self, wb) -> None:
        if self.QUEUE_SHEET not in wb.sheetnames:
            wb.create_sheet(self.QUEUE_SHEET).append(self.QUEUE_HEADER)
            return
        ws = wb[self.QUEUE_SHEET]
        existing = [ws.cell(row=1, column=c).value for c in range(1, len(self.QUEUE_HEADER) + 1)]
        if self._normalize_headers(existing) != self._normalize_headers(self.QUEUE_HEADER):
            raise ValueError("Queue sheet header mismatch")

    def enqueue_event(self, payload: dict, created_at: str, event_id: str | None = None) -> tuple[str, bool]:
        event_id = (event_id or str(uuid.uuid4())).strip()
        if not event_id:
            event_id = str(uuid.uuid4())
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb[self.QUEUE_SHEET]
            for row_idx in range(2, ws.max_row + 1):
                row_event = str(ws.cell(row=row_idx, column=1).value or "").strip()
                if row_event == event_id:
                    return event_id, False
            ws.append([event_id, "pending", 0, "", created_at, created_at, "", json.dumps(payload)])
            wb.save(self.path)
        return event_id, True

    def claim_due_events(self, now_iso_value: str, limit: int = 20) -> List[dict]:
        claimed: List[dict] = []
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb[self.QUEUE_SHEET]
            for row_idx in range(2, ws.max_row + 1):
                if len(claimed) >= limit:
                    break
                status = str(ws.cell(row=row_idx, column=2).value or "").strip().lower()
                next_retry_at = str(ws.cell(row=row_idx, column=4).value or "").strip()
                if status not in {"pending", "retry"}:
                    continue
                if next_retry_at and next_retry_at > now_iso_value:
                    continue
                attempt_count = int(ws.cell(row=row_idx, column=3).value or 0) + 1
                ws.cell(row=row_idx, column=2, value="processing")
                ws.cell(row=row_idx, column=3, value=attempt_count)
                ws.cell(row=row_idx, column=6, value=now_iso_value)
                raw_payload = str(ws.cell(row=row_idx, column=8).value or "{}")
                payload = json.loads(raw_payload)
                claimed.append(
                    {
                        "event_id": str(ws.cell(row=row_idx, column=1).value or "").strip(),
                        "attempt_count": attempt_count,
                        "payload": payload,
                    }
                )
            wb.save(self.path)
        return claimed

    def mark_event_processed(self, event_id: str, now_iso_value: str) -> None:
        self._update_event_status(event_id, "processed", now_iso_value, "", "")

    def mark_event_retry(self, event_id: str, now_iso_value: str, next_retry_at: str, error: str) -> None:
        self._update_event_status(event_id, "retry", now_iso_value, next_retry_at, error)

    def mark_event_failed(self, event_id: str, now_iso_value: str, error: str) -> None:
        self._update_event_status(event_id, "failed", now_iso_value, "", error)

    def _update_event_status(
        self, event_id: str, status: str, now_iso_value: str, next_retry_at: str, error: str
    ) -> None:
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb[self.QUEUE_SHEET]
            for row_idx in range(2, ws.max_row + 1):
                row_event = str(ws.cell(row=row_idx, column=1).value or "").strip()
                if row_event != event_id:
                    continue
                ws.cell(row=row_idx, column=2, value=status)
                ws.cell(row=row_idx, column=4, value=next_retry_at)
                ws.cell(row=row_idx, column=6, value=now_iso_value)
                ws.cell(row=row_idx, column=7, value=(error or "")[:1000])
                break
            wb.save(self.path)

    def get_queue_stats(self) -> dict:
        stats = {"pending": 0, "processing": 0, "processed": 0, "retry": 0, "failed": 0, "total": 0}
        with self._lock:
            wb = load_workbook(self.path, read_only=True)
            ws = wb[self.QUEUE_SHEET]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2:
                    continue
                status = str(row[1] or "").strip().lower()
                if status in stats:
                    stats[status] += 1
                stats["total"] += 1
        return stats

    def append_submission(self, record: SubmissionRecord) -> None:
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb[self.FEEDBACK_SHEET]
            ws.append(
                [
                    "submission",
                    record.employee_id,
                    record.employee_name,
                    record.year,
                    record.quarter,
                    record.form_type,
                    json.dumps(record.answers),
                    json.dumps(record.comments),
                    record.submitted_at,
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )
            wb.save(self.path)

    def get_employee_name(self, employee_id: str) -> str | None:
        with self._lock:
            wb = load_workbook(self.path, read_only=True)
            ws = wb[self.FEEDBACK_SHEET]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 3:
                    continue
                if str(row[1] or "").strip() == employee_id:
                    name = str(row[2] or "").strip()
                    if name:
                        return name
        return None

    def get_employee_quarter(self, employee_id: str, year: int, quarter: str) -> List[SubmissionRecord]:
        rows: List[SubmissionRecord] = []
        with self._lock:
            wb = load_workbook(self.path, read_only=True)
            ws = wb[self.FEEDBACK_SHEET]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 9:
                    continue
                record_type = str(row[0] or "").strip()
                if record_type != "submission":
                    continue
                if str(row[1] or "").strip() != employee_id:
                    continue
                if int(row[3]) != year:
                    continue
                if str(row[4] or "").strip() != quarter:
                    continue
                rows.append(self._row_to_record(row))
        return rows

    def list_submission_status(self, year: int, quarter: str) -> List[dict]:
        by_employee: dict[str, dict] = {}
        with self._lock:
            wb = load_workbook(self.path, read_only=True)
            ws = wb[self.FEEDBACK_SHEET]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 22:
                    continue
                record_type = str(row[0] or "").strip()
                row_emp = str(row[1] or "").strip()
                if not row_emp:
                    continue
                row_name = str(row[2] or "").strip() or row_emp
                row_year = str(row[3] or "").strip()
                row_quarter = str(row[4] or "").strip()
                if row_year != str(year) or row_quarter != quarter:
                    continue

                item = by_employee.setdefault(
                    row_emp,
                    {
                        "employee_email": row_emp,
                        "employee_name": row_name,
                        "submitted_forms": set(),
                        "manager_email": "",
                        "client_email": "",
                    },
                )

                if record_type == "submission":
                    form_type = str(row[5] or "").strip()
                    if form_type:
                        item["submitted_forms"].add(form_type)
                elif record_type == "summary":
                    item["manager_email"] = str(row[9] or "").strip()
                    item["client_email"] = str(row[10] or "").strip()

        results: List[dict] = []
        for _, item in by_employee.items():
            results.append(
                {
                    "employee_email": item["employee_email"],
                    "employee_name": item["employee_name"],
                    "submitted_forms": sorted(item["submitted_forms"]),
                    "manager_email": item["manager_email"],
                    "client_email": item["client_email"],
                }
            )
        return results

    def _row_to_record(self, row) -> SubmissionRecord:
        try:
            answers = [float(v) for v in json.loads(row[6] or "[]")]
        except (ValueError, TypeError, json.JSONDecodeError):
            answers = []
        try:
            comments = [str(v) for v in json.loads(row[7] or "[]")]
        except (TypeError, json.JSONDecodeError):
            comments = []
        return SubmissionRecord(
            employee_id=str(row[1] or "").strip(),
            employee_name=str(row[2] or "Unknown").strip(),
            year=int(row[3]),
            quarter=str(row[4] or "").strip(),
            form_type=str(row[5] or "").strip(),
            answers=answers,
            comments=comments,
            submitted_at=str(row[8] or "").strip(),
        )

    def upsert_summary_row(
        self,
        employee_id: str,
        employee_name: str,
        manager_email: str,
        client_email: str,
        peer_1: str,
        peer_2: str,
        self_avg: float | None,
        client_avg: float | None,
        peer_avg: float | None,
        manager_avg: float | None,
        total_average: float,
        client_diff_gt_1: float | str,
        manager_diff_gt_1: float | str,
        client_diff_gt_05: float | str,
        manager_diff_gt_05: float | str,
        year: int,
        quarter: str,
    ) -> None:
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb[self.FEEDBACK_SHEET]
            summary_ws = wb[self.summary_view_sheet]
            target_row = None
            for row_idx in range(2, ws.max_row + 1):
                if str(ws.cell(row=row_idx, column=1).value or "").strip() != "summary":
                    continue
                row_emp = str(ws.cell(row=row_idx, column=2).value or "").strip()
                row_year = str(ws.cell(row=row_idx, column=4).value or "").strip()
                row_quarter = str(ws.cell(row=row_idx, column=5).value or "").strip()
                if row_emp == employee_id and row_year == str(year) and row_quarter == quarter:
                    target_row = row_idx
                    break

            row_data = [
                "summary",
                employee_id,
                employee_name,
                year,
                quarter,
                "",
                "",
                "",
                now_iso(),
                manager_email or "",
                client_email or "",
                peer_1 or "",
                peer_2 or "",
                self_avg if self_avg is not None else "",
                client_avg if client_avg is not None else "",
                peer_avg if peer_avg is not None else "",
                manager_avg if manager_avg is not None else "",
                total_average,
                client_diff_gt_1,
                manager_diff_gt_1,
                client_diff_gt_05,
                manager_diff_gt_05,
            ]

            if target_row is None:
                ws.append(row_data)
            else:
                for col_idx, val in enumerate(row_data, start=1):
                    ws.cell(row=target_row, column=col_idx, value=val)

            # Mirror summary values into HR-style top table in summary view sheet.
            self._upsert_summary_view_row(
                ws=summary_ws,
                employee_id=employee_id,
                employee_name=employee_name,
                manager_email=manager_email,
                client_email=client_email,
                peer_1=peer_1,
                peer_2=peer_2,
                self_avg=self_avg,
                client_avg=client_avg,
                peer_avg=peer_avg,
                manager_avg=manager_avg,
                total_average=total_average,
                client_diff_gt_1=client_diff_gt_1,
                manager_diff_gt_1=manager_diff_gt_1,
                client_diff_gt_05=client_diff_gt_05,
                manager_diff_gt_05=manager_diff_gt_05,
            )
            wb.save(self.path)

    def _is_kpi_row(self, ws, row_idx: int) -> bool:
        name_cell = ws.cell(row=row_idx, column=2).value
        if name_cell not in (None, ""):
            return False
        # KPI block in sample sheet has values in columns 7-15 without employee identity fields.
        for col_idx in range(7, 16):
            if ws.cell(row=row_idx, column=col_idx).value not in (None, ""):
                return True
        label = str(ws.cell(row=row_idx, column=10).value or "").strip().lower()
        return label in {"total average", ">=4.5", ">3.5 and <4.5", "<=3.5 and >2", ">=4", ">3 and <4"}

    def _find_insert_row_before_kpi(self, ws) -> int:
        for row_idx in range(2, ws.max_row + 1):
            if self._is_kpi_row(ws, row_idx):
                return row_idx
        return ws.max_row + 1

    def _upsert_summary_view_row(
        self,
        ws,
        employee_id: str,
        employee_name: str,
        manager_email: str,
        client_email: str,
        peer_1: str,
        peer_2: str,
        self_avg: float | None,
        client_avg: float | None,
        peer_avg: float | None,
        manager_avg: float | None,
        total_average: float,
        client_diff_gt_1: float | str,
        manager_diff_gt_1: float | str,
        client_diff_gt_05: float | str,
        manager_diff_gt_05: float | str,
    ) -> None:
        target_row = None
        for row_idx in range(2, ws.max_row + 1):
            row_emp = str(ws.cell(row=row_idx, column=1).value or "").strip()
            row_name = str(ws.cell(row=row_idx, column=2).value or "").strip()
            if row_emp and row_emp == employee_id:
                target_row = row_idx
                break
            if not row_emp and row_name and row_name == employee_name:
                target_row = row_idx
                break

        row_data = [
            employee_id,
            employee_name,
            manager_email or "",
            client_email or "",
            peer_1 or "",
            peer_2 or "",
            self_avg if self_avg is not None else "",
            client_avg if client_avg is not None else "",
            peer_avg if peer_avg is not None else "",
            manager_avg if manager_avg is not None else "",
            total_average,
            client_diff_gt_1,
            manager_diff_gt_1,
            client_diff_gt_05,
            manager_diff_gt_05,
        ]

        if target_row is None:
            insert_at = self._find_insert_row_before_kpi(ws)
            ws.insert_rows(insert_at, 1)
            target_row = insert_at

        for col_idx, val in enumerate(row_data, start=1):
            ws.cell(row=target_row, column=col_idx, value=val)

        self._refresh_kpi_formulas(ws)

    def _find_row_with_label(self, ws, col_idx: int, label: str) -> int | None:
        for row_idx in range(2, ws.max_row + 1):
            value = str(ws.cell(row=row_idx, column=col_idx).value or "").strip().lower()
            if value == label.strip().lower():
                return row_idx
        return None

    def _refresh_kpi_formulas(self, ws) -> None:
        kpi_start = self._find_insert_row_before_kpi(ws)
        data_start = 2
        data_end = kpi_start - 1
        if data_end < data_start:
            return

        b_rng = f"$B${data_start}:$B${data_end}"  # name present indicates employee row
        h_rng = f"$H${data_start}:$H${data_end}"  # client_avg
        j_rng = f"$J${data_start}:$J${data_end}"  # manager_avg
        k_rng = f"$K${data_start}:$K${data_end}"  # total_average
        l_rng = f"$L${data_start}:$L${data_end}"  # client diff >1 marker
        m_rng = f"$M${data_start}:$M${data_end}"  # manager diff >1 marker
        n_rng = f"$N${data_start}:$N${data_end}"  # client diff >0.5 marker
        o_rng = f"$O${data_start}:$O${data_end}"  # manager diff >0.5 marker

        # Total Average distribution table (column J labels, column K values).
        avg_ge_45 = self._find_row_with_label(ws, 10, ">=4.5")
        if avg_ge_45:
            ws.cell(
                row=avg_ge_45,
                column=11,
                value=f'=COUNTIFS({b_rng},"<>",{k_rng},">=4.5")',
            )
        avg_gt_35_lt_45 = self._find_row_with_label(ws, 10, ">3.5 and <4.5")
        if avg_gt_35_lt_45:
            ws.cell(
                row=avg_gt_35_lt_45,
                column=11,
                value=f'=COUNTIFS({b_rng},"<>",{k_rng},">3.5",{k_rng},"<4.5")',
            )
        avg_le_35_gt_2 = self._find_row_with_label(ws, 10, "<=3.5 and >2")
        if avg_le_35_gt_2:
            ws.cell(
                row=avg_le_35_gt_2,
                column=11,
                value=f'=COUNTIFS({b_rng},"<>",{k_rng},"<=3.5",{k_rng},">2")',
            )

        # Difference percentage table (column M labels, columns N/O values as percentages).
        client_diff_row = self._find_row_with_label(ws, 13, "Client vs Self Difference")
        if client_diff_row:
            ws.cell(
                row=client_diff_row,
                column=14,
                value=f'=IFERROR(COUNTIFS({b_rng},"<>",{l_rng},"<>")/COUNTIFS({b_rng},"<>",{h_rng},"<>")*100,0)',
            )
            ws.cell(
                row=client_diff_row,
                column=15,
                value=f'=IFERROR(COUNTIFS({b_rng},"<>",{n_rng},"<>")/COUNTIFS({b_rng},"<>",{h_rng},"<>")*100,0)',
            )

        manager_diff_row = self._find_row_with_label(ws, 13, "Manager vs Self Difference")
        if manager_diff_row:
            ws.cell(
                row=manager_diff_row,
                column=14,
                value=f'=IFERROR(COUNTIFS({b_rng},"<>",{m_rng},"<>")/COUNTIFS({b_rng},"<>",{j_rng},"<>")*100,0)',
            )
            ws.cell(
                row=manager_diff_row,
                column=15,
                value=f'=IFERROR(COUNTIFS({b_rng},"<>",{o_rng},"<>")/COUNTIFS({b_rng},"<>",{j_rng},"<>")*100,0)',
            )


def now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"
